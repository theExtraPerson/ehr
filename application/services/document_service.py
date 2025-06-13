from docx import Document
from io import BytesIO
from datetime import datetime
import os

class DocumentService:
    TEMPLATE_PATH = os.path.join(os.path.dirname(__file__), '../templates/docs')
    
    @classmethod
    def generate_prescription(cls, prescription):
        """Generate prescription as Word doc"""
        doc = Document(os.path.join(cls.TEMPLATE_PATH, 'prescription_template.docx'))
        
        # Replace placeholders
        for paragraph in doc.paragraphs:
            paragraph.text = paragraph.text.replace('[PATIENT_NAME]', prescription.patient.full_name())
            paragraph.text = paragraph.text.replace('[MEDICATION]', prescription.medication)
            paragraph.text = paragraph.text.replace('[DOSAGE]', prescription.dosage)
            paragraph.text = paragraph.text.replace('[DATE]', datetime.now().strftime('%Y-%m-%d'))
        
        # Save to memory
        output = BytesIO()
        doc.save(output)
        output.seek(0)
        return output
    
    @classmethod
    def generate_pdf_invoice(cls, invoice):
        """Generate PDF invoice as fallback"""
        from reportlab.pdfgen import canvas
        from reportlab.lib.pagesizes import letter
        
        buffer = BytesIO()
        p = canvas.Canvas(buffer, pagesize=letter)
        
        # Draw invoice
        p.setFont("Helvetica-Bold", 16)
        p.drawString(100, 750, f"Invoice #{invoice.id}")
        
        p.setFont("Helvetica", 12)
        p.drawString(100, 720, f"Patient: {invoice.patient.full_name()}")
        p.drawString(100, 700, f"Date: {invoice.created_at.strftime('%Y-%m-%d')}")
        p.drawString(100, 680, f"Amount Due: ${invoice.total_amount:,.2f}")
        
        p.showPage()
        p.save()
        buffer.seek(0)
        return buffer